"""
Store Invoice & Purchase Order Tracker
=====================================

Flask + SQLite backend. The frontend (static/app.js) talks to the JSON API
defined below.

Run:  store-invoice          (after pip install)
      python app.py          (from the source folder)
      ->   http://127.0.0.1:5000
"""

import io
import logging
import os
import re
import sqlite3
from datetime import date, datetime, timedelta
from decimal import ROUND_HALF_UP, Decimal, InvalidOperation

from flask import Flask, g, jsonify, render_template, request, send_file
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.datavalidation import DataValidation
from werkzeug.exceptions import HTTPException

# ---------------------------------------------------------------------------
# Configuration
# ---------------------------------------------------------------------------

def default_db_path():
    """STORE_INVOICE_DB if set, else a per-user data folder (safe for pip installs)."""
    env = os.environ.get("STORE_INVOICE_DB")
    if env:
        return os.path.abspath(env)
    base = (os.environ.get("LOCALAPPDATA") or os.environ.get("XDG_DATA_HOME")
            or os.path.join(os.path.expanduser("~"), ".local", "share"))
    folder = os.path.join(base, "StoreInvoice")
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, "store_invoice.db")


DB_PATH = default_db_path()
MAX_UPLOAD_MB = int(os.environ.get("MAX_UPLOAD_MB", "10"))

INVOICE_TYPES = ["Purchase", "Service", "Other"]
PAYMENT_STATUSES = ["Pending", "Paid", "Partially Paid", "Hold"]
RECEIVED_OPTIONS = ["Yes", "No"]
PO_STATUSES = ["Matched", "Rate Mismatch", "PO Pending", "Check Rate"]
COMMON_UNITS = ["KG", "GM", "QTL", "TON", "LTR", "ML", "NOS", "PCS", "BAG", "BOX",
                "PKT", "CTN", "DOZ", "SET", "MTR", "ROLL", "BTL", "DRUM"]

# Values closer than these are treated as equal when matching against the PO.
RATE_TOLERANCE = 0.005    # rupees per unit
AMOUNT_TOLERANCE = 1.00   # rupees per line (absorbs rounding)

MAX_ITEMS_PER_INVOICE = 500
MAX_TEXT = {"invoice_no": 60, "vendor_name": 150, "remarks": 1000, "particulars": 250,
            "hsn": 20, "unit": 20, "po_number": 60}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_MB * 1024 * 1024
app.json.sort_keys = False

logging.basicConfig(level=logging.INFO, format="%(asctime)s %(levelname)s %(message)s")
log = logging.getLogger("store_invoice")


class ApiError(Exception):
    """An error whose message is safe to show to the user."""

    def __init__(self, message, status=400, details=None):
        super().__init__(message)
        self.message = message
        self.status = status
        self.details = details or []


# ---------------------------------------------------------------------------
# Database
# ---------------------------------------------------------------------------

SCHEMA = """
CREATE TABLE IF NOT EXISTS invoices (
    id               INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_date     TEXT NOT NULL,
    invoice_no       TEXT NOT NULL,
    vendor_name      TEXT NOT NULL,
    invoice_type     TEXT NOT NULL DEFAULT 'Purchase',
    payment_status   TEXT NOT NULL DEFAULT 'Pending',
    invoice_received TEXT NOT NULL DEFAULT 'Yes',
    remarks          TEXT NOT NULL DEFAULT '',
    created_at       TEXT NOT NULL,
    updated_at       TEXT NOT NULL
);

CREATE TABLE IF NOT EXISTS invoice_items (
    id          INTEGER PRIMARY KEY AUTOINCREMENT,
    invoice_id  INTEGER NOT NULL REFERENCES invoices(id) ON DELETE CASCADE,
    particulars TEXT NOT NULL,
    hsn         TEXT NOT NULL DEFAULT '',
    rate        REAL,
    quantity    REAL,
    unit        TEXT NOT NULL DEFAULT '',
    gst         REAL NOT NULL DEFAULT 0,
    without_gst REAL,
    amount      REAL,
    po_date     TEXT,
    po_number   TEXT NOT NULL DEFAULT '',
    po_rate     REAL,
    po_amount   REAL
);

CREATE INDEX IF NOT EXISTS idx_invoices_date   ON invoices(invoice_date);
CREATE INDEX IF NOT EXISTS idx_invoices_vendor ON invoices(vendor_name);
CREATE INDEX IF NOT EXISTS idx_invoices_no     ON invoices(invoice_no);
CREATE INDEX IF NOT EXISTS idx_items_invoice   ON invoice_items(invoice_id);
CREATE INDEX IF NOT EXISTS idx_items_po        ON invoice_items(po_number);
"""


def connect_db():
    conn = sqlite3.connect(DB_PATH, timeout=15)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = connect_db()
    try:
        conn.executescript(SCHEMA)
        conn.commit()
    finally:
        conn.close()
    log.info("SQLite database ready at %s", DB_PATH)


def get_db():
    if "db" not in g:
        g.db = connect_db()
    return g.db


@app.teardown_appcontext
def close_db(_exc):
    db = g.pop("db", None)
    if db is not None:
        db.close()


def now_ts():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------------------
# Parsing helpers
# ---------------------------------------------------------------------------

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], start=1)}


def round2(value):
    if value is None:
        return None
    return float(Decimal(str(value)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def parse_number(value):
    """Return a float, None for blank, or raise ValueError for junk."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("not a number")
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"(?i)^(rs\.?|inr)\s*", "", text).replace("\u20b9", "")
    text = text.replace(",", "").replace("%", "").strip()
    try:
        number = Decimal(text)
    except InvalidOperation:
        raise ValueError("not a number")
    if not number.is_finite():
        raise ValueError("not a number")
    return float(number)


def _make_date(year, month, day):
    if year < 100:
        year += 2000
    try:
        return date(year, month, day)
    except ValueError:
        return None


def parse_date(value):
    """Accepts date/datetime, Excel serial numbers and common text formats.

    Text dates are read day-first (Indian convention): 01-10-2026 = 1 Oct 2026.
    Returns an ISO string (YYYY-MM-DD), None for blank, raises ValueError otherwise.
    """
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if 1 <= value < 2958466:
            return from_excel(value).date().isoformat()
        raise ValueError("bad date")
    text = str(value).strip()
    if not text:
        return None
    text = re.sub(r"[T ]\d{1,2}:\d{2}(:\d{2})?(\.\d+)?$", "", text)  # drop a time part
    d = None
    m = re.fullmatch(r"(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})", text)
    if m:
        d = _make_date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    if d is None:
        m = re.fullmatch(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{2}|\d{4})", text)
        if m:
            d = _make_date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    if d is None:
        m = re.fullmatch(r"(\d{1,2})(?:st|nd|rd|th)?[\s\-/.,]*([A-Za-z]{3,9})[\s\-/.,]*(\d{2}|\d{4})", text)
        if m and m.group(2)[:3].lower() in MONTHS:
            d = _make_date(int(m.group(3)), MONTHS[m.group(2)[:3].lower()], int(m.group(1)))
    if d is None:
        m = re.fullmatch(r"([A-Za-z]{3,9})[\s\-/.]*(\d{1,2})(?:st|nd|rd|th)?[\s,\-/.]*(\d{4})", text)
        if m and m.group(1)[:3].lower() in MONTHS:
            d = _make_date(int(m.group(3)), MONTHS[m.group(1)[:3].lower()], int(m.group(2)))
    if d is None and re.fullmatch(r"\d{5}(\.\d+)?", text):
        return parse_date(float(text))
    if d is None:
        raise ValueError("bad date")
    return d.isoformat()


def clean_text(value, max_len=None):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)  # e.g. HSN 1101 read from Excel as 1101.0
    text = re.sub(r"\s+", " ", str(value)).strip()
    if max_len:
        text = text[:max_len]
    return text


def pick_option(value, options, default, aliases=None):
    """Case-insensitive match of value to one of options; blank -> default."""
    text = clean_text(value).lower()
    if not text:
        return default
    for opt in options:
        if opt.lower() == text:
            return opt
    if aliases and text in aliases:
        return aliases[text]
    return None


PAYMENT_ALIASES = {"partial": "Partially Paid", "partly paid": "Partially Paid",
                   "part paid": "Partially Paid", "unpaid": "Pending", "on hold": "Hold"}
RECEIVED_ALIASES = {"y": "Yes", "n": "No", "true": "Yes", "false": "No",
                    "received": "Yes", "not received": "No", "1": "Yes", "0": "No"}


# ---------------------------------------------------------------------------
# Calculations & PO matching
# ---------------------------------------------------------------------------

def compute_item(item):
    """Adds calculated fields to an item dict (rate/quantity/gst already numeric)."""
    rate, qty, gst = item.get("rate"), item.get("quantity"), item.get("gst") or 0.0
    if rate is not None and qty is not None:
        without_gst = round2(rate * qty)
        amount = round2(rate * qty * (1 + gst / 100.0))
    else:
        without_gst = amount = None
    item["gst"] = gst
    item["without_gst"] = without_gst
    item["amount"] = amount
    item["gst_amount"] = round2(amount - without_gst) if amount is not None else None
    item["rate_diff"] = (round(rate - item["po_rate"], 4)
                         if rate is not None and item.get("po_rate") is not None else None)
    item["amount_diff"] = (round2(amount - item["po_amount"])
                           if amount is not None and item.get("po_amount") is not None else None)
    item["po_status"], item["po_status_reason"] = item_po_status(item)
    return item


def item_po_status(it):
    """Returns (status, human readable reason)."""
    rate, qty, amount = it.get("rate"), it.get("quantity"), it.get("amount")
    if not it.get("particulars") or rate is None or qty is None:
        return "Check Rate", "Particulars, rate or quantity is missing"
    if not (it.get("po_number") or "").strip() or it.get("po_rate") is None:
        return "PO Pending", "PO number or PO rate not entered yet"
    if abs(rate - it["po_rate"]) > RATE_TOLERANCE:
        return "Rate Mismatch", "Invoice rate differs from PO rate by %s" % _fmt(rate - it["po_rate"])
    po_amount = it.get("po_amount")
    if po_amount is None:
        return "Check Rate", "Rate matches but PO amount is missing"
    if abs(amount - po_amount) <= AMOUNT_TOLERANCE:
        return "Matched", "Rate and amount match the PO"
    if abs(it["without_gst"] - po_amount) <= AMOUNT_TOLERANCE:
        return "Matched", "Rate matches; PO amount matches the value before GST"
    return "Check Rate", "Rate matches but amount differs from PO amount by %s" % _fmt(amount - po_amount)


def _fmt(n):
    return "{:+,.2f}".format(n)


def invoice_po_status(item_statuses):
    if not item_statuses:
        return "Check Rate"
    for status in ("Rate Mismatch", "PO Pending", "Check Rate"):
        if status in item_statuses:
            return status
    return "Matched"


def invoice_totals(items):
    def total(key):
        return round2(sum(i[key] for i in items if i.get(key) is not None))
    return {
        "item_count": len(items),
        "quantity": round(sum(i["quantity"] or 0 for i in items), 3),
        "without_gst": total("without_gst"),
        "gst_amount": total("gst_amount"),
        "amount": total("amount"),
        "po_amount": total("po_amount"),
        "difference": total("amount_diff"),
    }


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def validate_invoice(data):
    """Validate a JSON payload / imported group. Returns (header, items, errors)."""
    errors = []
    if not isinstance(data, dict):
        return None, None, ["Invalid request data."]

    header = {}
    try:
        header["invoice_date"] = parse_date(data.get("invoice_date"))
        if not header["invoice_date"]:
            errors.append("Invoice Date is required.")
    except ValueError:
        errors.append("Invoice Date is not a valid date (use DD-MM-YYYY).")

    for key, label in (("invoice_no", "Invoice No."), ("vendor_name", "Vendor Name")):
        header[key] = clean_text(data.get(key), MAX_TEXT[key])
        if not header[key]:
            errors.append("%s is required." % label)

    header["invoice_type"] = pick_option(data.get("invoice_type"), INVOICE_TYPES, "Purchase")
    if header["invoice_type"] is None:
        errors.append("Invoice Type must be one of: %s." % ", ".join(INVOICE_TYPES))
    header["payment_status"] = pick_option(data.get("payment_status"), PAYMENT_STATUSES, "Pending",
                                           PAYMENT_ALIASES)
    if header["payment_status"] is None:
        errors.append("Payment Status must be one of: %s." % ", ".join(PAYMENT_STATUSES))
    header["invoice_received"] = pick_option(data.get("invoice_received"), RECEIVED_OPTIONS, "Yes",
                                             RECEIVED_ALIASES)
    if header["invoice_received"] is None:
        errors.append("Invoice Received must be Yes or No.")
    header["remarks"] = clean_text(data.get("remarks"), MAX_TEXT["remarks"])

    raw_items = data.get("items")
    items = []
    if not isinstance(raw_items, list) or not raw_items:
        errors.append("Add at least one item.")
        raw_items = []
    if len(raw_items) > MAX_ITEMS_PER_INVOICE:
        errors.append("An invoice can have at most %d items." % MAX_ITEMS_PER_INVOICE)
        raw_items = []

    for idx, raw in enumerate(raw_items, start=1):
        label = raw.get("_label") if isinstance(raw, dict) else None
        label = label or "Item %d" % idx
        if not isinstance(raw, dict):
            errors.append("%s: invalid data." % label)
            continue
        item = {
            "particulars": clean_text(raw.get("particulars"), MAX_TEXT["particulars"]),
            "hsn": clean_text(raw.get("hsn"), MAX_TEXT["hsn"]),
            "unit": clean_text(raw.get("unit"), MAX_TEXT["unit"]).upper(),
            "po_number": clean_text(raw.get("po_number"), MAX_TEXT["po_number"]),
        }
        if not item["particulars"]:
            errors.append("%s: Particulars is required." % label)

        for key, name, required in (("rate", "Rate", True), ("quantity", "Quantity", True),
                                    ("gst", "GST %", False), ("po_rate", "PO Rate", False),
                                    ("po_amount", "PO Amount", False)):
            try:
                val = parse_number(raw.get(key))
            except ValueError:
                errors.append("%s: %s must be a number." % (label, name))
                val = None
            else:
                if val is None and required:
                    errors.append("%s: %s is required." % (label, name))
                elif val is not None and val < 0:
                    errors.append("%s: %s cannot be negative." % (label, name))
                    val = None
            item[key] = val
        if item["gst"] is not None and item["gst"] > 100:
            errors.append("%s: GST %% cannot be more than 100." % label)
        try:
            item["po_date"] = parse_date(raw.get("po_date"))
        except ValueError:
            errors.append("%s: PO Date is not a valid date." % label)
            item["po_date"] = None
        items.append(item)

    if not errors:
        items = [compute_item(i) for i in items]
    return header, items, errors


# ---------------------------------------------------------------------------
# Queries
# ---------------------------------------------------------------------------

FILTER_KEYS = ("search", "start_date", "end_date", "vendor", "payment_status",
               "po_status", "invoice_received", "invoice_type")


def read_filters(args):
    filters = {k: (args.get(k) or "").strip() for k in FILTER_KEYS}
    for key, label in (("start_date", "Start Date"), ("end_date", "End Date")):
        if filters[key]:
            try:
                filters[key] = parse_date(filters[key])
            except ValueError:
                raise ApiError("%s is not a valid date." % label)
    if filters["start_date"] and filters["end_date"] and filters["start_date"] > filters["end_date"]:
        raise ApiError("Start Date cannot be after End Date.")
    if filters["po_status"] and filters["po_status"] not in PO_STATUSES:
        raise ApiError("Unknown PO status filter.")
    return filters


def _like(term):
    escaped = term.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")
    return "%" + escaped + "%"


def fetch_invoices(db, filters):
    clauses, params = [], []
    if filters.get("start_date"):
        clauses.append("i.invoice_date >= ?")
        params.append(filters["start_date"])
    if filters.get("end_date"):
        clauses.append("i.invoice_date <= ?")
        params.append(filters["end_date"])
    if filters.get("vendor"):
        clauses.append("LOWER(i.vendor_name) = LOWER(?)")
        params.append(filters["vendor"])
    for key in ("payment_status", "invoice_received", "invoice_type"):
        if filters.get(key):
            clauses.append("i.%s = ?" % key)  # key comes from a fixed whitelist
            params.append(filters[key])
    if filters.get("search"):
        term = _like(filters["search"])
        clauses.append(
            "(i.invoice_no LIKE ? ESCAPE '\\' OR i.vendor_name LIKE ? ESCAPE '\\' OR EXISTS ("
            " SELECT 1 FROM invoice_items s WHERE s.invoice_id = i.id AND ("
            " s.particulars LIKE ? ESCAPE '\\' OR s.po_number LIKE ? ESCAPE '\\'"
            " OR s.hsn LIKE ? ESCAPE '\\')))")
        params.extend([term] * 5)
    sql = "SELECT i.* FROM invoices i"
    if clauses:
        sql += " WHERE " + " AND ".join(clauses)
    sql += " ORDER BY i.invoice_date DESC, i.id DESC"
    rows = db.execute(sql, params).fetchall()

    items_by_invoice = load_items(db, [r["id"] for r in rows])
    invoices = [build_invoice(r, items_by_invoice.get(r["id"], [])) for r in rows]
    if filters.get("po_status"):
        invoices = [inv for inv in invoices if inv["po_status"] == filters["po_status"]]
    for n, inv in enumerate(invoices, start=1):
        inv["sl_no"] = n
    return invoices


def load_items(db, invoice_ids):
    result = {}
    for start in range(0, len(invoice_ids), 500):
        chunk = invoice_ids[start:start + 500]
        placeholders = ",".join("?" * len(chunk))
        for row in db.execute(
                "SELECT * FROM invoice_items WHERE invoice_id IN (%s) ORDER BY id" % placeholders,
                chunk):
            item = dict(row)
            result.setdefault(item["invoice_id"], []).append(compute_item(item))
    return result


def build_invoice(row, items):
    inv = dict(row)
    inv["items"] = items
    inv["totals"] = invoice_totals(items)
    inv["po_status"] = invoice_po_status([i["po_status"] for i in items])
    return inv


def get_invoice(db, invoice_id):
    row = db.execute("SELECT * FROM invoices WHERE id = ?", (invoice_id,)).fetchone()
    if row is None:
        return None
    return build_invoice(row, load_items(db, [invoice_id]).get(invoice_id, []))


def find_duplicate(db, header, exclude_id=None):
    sql = ("SELECT id FROM invoices WHERE LOWER(invoice_no) = LOWER(?) "
           "AND LOWER(vendor_name) = LOWER(?)")
    params = [header["invoice_no"], header["vendor_name"]]
    if exclude_id is not None:
        sql += " AND id != ?"
        params.append(exclude_id)
    row = db.execute(sql, params).fetchone()
    return row["id"] if row else None


ITEM_COLUMNS = ("particulars", "hsn", "rate", "quantity", "unit", "gst", "without_gst",
                "amount", "po_date", "po_number", "po_rate", "po_amount")


def insert_items(db, invoice_id, items):
    db.executemany(
        "INSERT INTO invoice_items (invoice_id, %s) VALUES (?, %s)"
        % (", ".join(ITEM_COLUMNS), ", ".join("?" * len(ITEM_COLUMNS))),
        [[invoice_id] + [it[c] for c in ITEM_COLUMNS] for it in items])


def insert_invoice(db, header, items):
    ts = now_ts()
    cur = db.execute(
        "INSERT INTO invoices (invoice_date, invoice_no, vendor_name, invoice_type, payment_status,"
        " invoice_received, remarks, created_at, updated_at) VALUES (?,?,?,?,?,?,?,?,?)",
        (header["invoice_date"], header["invoice_no"], header["vendor_name"], header["invoice_type"],
         header["payment_status"], header["invoice_received"], header["remarks"], ts, ts))
    insert_items(db, cur.lastrowid, items)
    return cur.lastrowid


def update_invoice(db, invoice_id, header, items):
    db.execute(
        "UPDATE invoices SET invoice_date=?, invoice_no=?, vendor_name=?, invoice_type=?,"
        " payment_status=?, invoice_received=?, remarks=?, updated_at=? WHERE id=?",
        (header["invoice_date"], header["invoice_no"], header["vendor_name"], header["invoice_type"],
         header["payment_status"], header["invoice_received"], header["remarks"], now_ts(),
         invoice_id))
    db.execute("DELETE FROM invoice_items WHERE invoice_id = ?", (invoice_id,))
    insert_items(db, invoice_id, items)


# ---------------------------------------------------------------------------
# Error handling
# ---------------------------------------------------------------------------

@app.errorhandler(ApiError)
def handle_api_error(err):
    body = {"error": err.message}
    if err.details:
        body["details"] = err.details
    return jsonify(body), err.status


@app.errorhandler(HTTPException)
def handle_http_error(err):
    messages = {
        404: "Not found.",
        405: "This action is not allowed.",
        413: "The file is too large. Maximum upload size is %d MB." % MAX_UPLOAD_MB,
    }
    if request.path.startswith("/api/") or err.code == 413:
        return jsonify({"error": messages.get(err.code, err.description or err.name)}), err.code
    return err


@app.errorhandler(sqlite3.Error)
def handle_db_error(err):
    log.exception("Database error: %s", err)
    return jsonify({"error": "A database error occurred. Please try again."}), 500


@app.errorhandler(Exception)
def handle_unexpected(err):
    log.exception("Unhandled error: %s", err)
    return jsonify({"error": "Something went wrong on the server. Please try again."}), 500


def json_body():
    data = request.get_json(silent=True)
    if data is None:
        raise ApiError("Request body must be JSON.")
    return data


# ---------------------------------------------------------------------------
# Pages & meta
# ---------------------------------------------------------------------------

@app.route("/")
def index():
    return render_template("index.html")


@app.route("/favicon.ico")
def favicon():
    return "", 204


@app.get("/api/meta")
def api_meta():
    return jsonify({
        "invoice_types": INVOICE_TYPES,
        "payment_statuses": PAYMENT_STATUSES,
        "received_options": RECEIVED_OPTIONS,
        "po_statuses": PO_STATUSES,
        "units": COMMON_UNITS,
        "rate_tolerance": RATE_TOLERANCE,
        "amount_tolerance": AMOUNT_TOLERANCE,
        "max_upload_mb": MAX_UPLOAD_MB,
        "ocr_provider": os.environ.get("OCR_PROVIDER", "tesseract"),
    })


@app.get("/api/vendors")
def api_vendors():
    rows = get_db().execute(
        "SELECT vendor_name, COUNT(*) AS n FROM invoices GROUP BY LOWER(vendor_name) "
        "ORDER BY vendor_name COLLATE NOCASE").fetchall()
    return jsonify({"vendors": [r["vendor_name"] for r in rows]})


# ---------------------------------------------------------------------------
# Invoice API
# ---------------------------------------------------------------------------

@app.get("/api/invoices")
def api_list_invoices():
    filters = read_filters(request.args)
    invoices = fetch_invoices(get_db(), filters)
    summaries = []
    for inv in invoices:
        s = {k: v for k, v in inv.items() if k != "items"}
        s["status_counts"] = {st: sum(1 for i in inv["items"] if i["po_status"] == st)
                              for st in PO_STATUSES}
        summaries.append(s)
    return jsonify({
        "invoices": summaries,
        "count": len(summaries),
        "totals": {
            "amount": round2(sum(i["totals"]["amount"] for i in invoices)),
            "po_amount": round2(sum(i["totals"]["po_amount"] for i in invoices)),
            "difference": round2(sum(i["totals"]["difference"] for i in invoices)),
        },
        "filters": filters,
    })


@app.get("/api/invoices/<int:invoice_id>")
def api_get_invoice(invoice_id):
    inv = get_invoice(get_db(), invoice_id)
    if inv is None:
        raise ApiError("Invoice not found.", 404)
    return jsonify(inv)


@app.post("/api/invoices")
def api_create_invoice():
    header, items, errors = validate_invoice(json_body())
    if errors:
        raise ApiError("Please correct the highlighted problems.", 400, errors)
    db = get_db()
    if find_duplicate(db, header):
        raise ApiError("Invoice %s from %s already exists." % (header["invoice_no"], header["vendor_name"]),
                       409)
    with db:  # one transaction for the invoice and all its items
        invoice_id = insert_invoice(db, header, items)
    return jsonify(get_invoice(db, invoice_id)), 201


@app.put("/api/invoices/<int:invoice_id>")
def api_update_invoice(invoice_id):
    db = get_db()
    if not db.execute("SELECT 1 FROM invoices WHERE id = ?", (invoice_id,)).fetchone():
        raise ApiError("Invoice not found.", 404)
    header, items, errors = validate_invoice(json_body())
    if errors:
        raise ApiError("Please correct the highlighted problems.", 400, errors)
    if find_duplicate(db, header, exclude_id=invoice_id):
        raise ApiError("Another invoice %s from %s already exists." % (header["invoice_no"],
                                                                      header["vendor_name"]), 409)
    with db:
        update_invoice(db, invoice_id, header, items)
    return jsonify(get_invoice(db, invoice_id))


@app.delete("/api/invoices/<int:invoice_id>")
def api_delete_invoice(invoice_id):
    db = get_db()
    with db:
        cur = db.execute("DELETE FROM invoices WHERE id = ?", (invoice_id,))
    if cur.rowcount == 0:
        raise ApiError("Invoice not found.", 404)
    return jsonify({"message": "Invoice deleted.", "id": invoice_id})


# ---------------------------------------------------------------------------
# Dashboard
# ---------------------------------------------------------------------------

@app.get("/api/dashboard")
def api_dashboard():
    filters = read_filters(request.args)
    invoices = fetch_invoices(get_db(), filters)
    items = [it for inv in invoices for it in inv["items"]]
    status_counts = {st: sum(1 for i in items if i["po_status"] == st) for st in PO_STATUSES}

    daily, vendors = {}, {}
    for inv in invoices:
        t = inv["totals"]
        d = daily.setdefault(inv["invoice_date"], {"date": inv["invoice_date"], "invoice_count": 0,
                                                   "invoice_value": 0.0, "po_value": 0.0,
                                                   "difference": 0.0, "item_count": 0})
        d["invoice_count"] += 1
        d["item_count"] += t["item_count"]
        d["invoice_value"] += t["amount"]
        d["po_value"] += t["po_amount"]
        d["difference"] += t["difference"]
        v = vendors.setdefault(inv["vendor_name"].lower(), {"vendor": inv["vendor_name"],
                                                            "invoice_count": 0, "invoice_value": 0.0})
        v["invoice_count"] += 1
        v["invoice_value"] += t["amount"]
    for d in daily.values():
        for k in ("invoice_value", "po_value", "difference"):
            d[k] = round2(d[k])
    for v in vendors.values():
        v["invoice_value"] = round2(v["invoice_value"])

    return jsonify({
        "filters": filters,
        "summary": {
            "total_invoices": len(invoices),
            "total_items": len(items),
            "total_invoice_value": round2(sum(i["totals"]["amount"] for i in invoices)),
            "total_without_gst": round2(sum(i["totals"]["without_gst"] for i in invoices)),
            "total_gst": round2(sum(i["totals"]["gst_amount"] for i in invoices)),
            "total_po_value": round2(sum(i["totals"]["po_amount"] for i in invoices)),
            "total_difference": round2(sum(i["totals"]["difference"] for i in invoices)),
            "matched_items": status_counts["Matched"],
            "rate_mismatch_items": status_counts["Rate Mismatch"],
            "po_pending_items": status_counts["PO Pending"],
            "check_rate_items": status_counts["Check Rate"],
            "total_quantity": round(sum(i["quantity"] or 0 for i in items), 3),
            "payment_pending": sum(1 for i in invoices if i["payment_status"] != "Paid"),
            "not_received": sum(1 for i in invoices if i["invoice_received"] == "No"),
        },
        "status_counts": status_counts,
        "daily": sorted(daily.values(), key=lambda d: d["date"]),
        "top_vendors": sorted(vendors.values(), key=lambda v: -v["invoice_value"])[:8],
    })


# ---------------------------------------------------------------------------
# Excel: export, template, import
# ---------------------------------------------------------------------------

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="D0D7DE")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
STATUS_FILLS = {
    "Matched": PatternFill("solid", fgColor="D4EDDA"),
    "Rate Mismatch": PatternFill("solid", fgColor="F8D7DA"),
    "PO Pending": PatternFill("solid", fgColor="FFF3CD"),
    "Check Rate": PatternFill("solid", fgColor="E2E3E5"),
}
FMT_MONEY = "#,##0.00"
FMT_RATE = "#,##0.00##"
FMT_QTY = "#,##0.###"
FMT_DATE = "DD-MM-YYYY"

# (header, key, number format)
EXPORT_COLUMNS = [
    ("SL No.", "sl_no", "0"),
    ("Invoice Date", "invoice_date", FMT_DATE),
    ("Invoice No.", "invoice_no", None),
    ("Vendor Name", "vendor_name", None),
    ("Invoice Type", "invoice_type", None),
    ("Payment Status", "payment_status", None),
    ("Invoice Received", "invoice_received", None),
    ("Remarks", "remarks", None),
    ("Particulars", "particulars", None),
    ("HSN", "hsn", "@"),
    ("Rate", "rate", FMT_RATE),
    ("Quantity", "quantity", FMT_QTY),
    ("Unit", "unit", None),
    ("GST %", "gst", "0.##"),
    ("Without GST", "without_gst", FMT_MONEY),
    ("Amount", "amount", FMT_MONEY),
    ("PO Date", "po_date", FMT_DATE),
    ("PO Number", "po_number", None),
    ("PO Rate", "po_rate", FMT_RATE),
    ("PO Amount", "po_amount", FMT_MONEY),
    ("PO Rate Difference", "rate_diff", FMT_RATE),
    ("PO Amount Difference", "amount_diff", FMT_MONEY),
    ("PO Match Status", "po_status", None),
]

EXPORT_HEADER_KEYS = {"sl_no", "invoice_date", "invoice_no", "vendor_name", "invoice_type",
                      "payment_status", "invoice_received", "remarks"}

TEMPLATE_COLUMNS = [
    ("Invoice Date", FMT_DATE, 14), ("Invoice No.", None, 14), ("Vendor Name", None, 26),
    ("Invoice Type", None, 13), ("Payment Status", None, 15), ("Invoice Received", None, 16),
    ("Remarks", None, 22), ("Particulars", None, 22), ("HSN", "@", 10), ("Rate", FMT_RATE, 10),
    ("Quantity", FMT_QTY, 10), ("Unit", None, 8), ("GST %", "0.##", 8), ("PO Date", FMT_DATE, 13),
    ("PO Number", None, 12), ("PO Rate", FMT_RATE, 10), ("PO Amount", FMT_MONEY, 13),
]

IMPORT_ALIASES = {
    "invoice_date": ["invoice date", "inv date", "bill date", "date"],
    "invoice_no": ["invoice no", "invoice number", "inv no", "invoice", "bill no", "bill number"],
    "vendor_name": ["vendor name", "vendor", "supplier", "supplier name", "party", "party name"],
    "invoice_type": ["invoice type", "type"],
    "payment_status": ["payment status", "payment"],
    "invoice_received": ["invoice received", "received"],
    "remarks": ["remarks", "remark", "notes", "note"],
    "particulars": ["particulars", "particular", "item", "item name", "description", "material"],
    "hsn": ["hsn", "hsn code", "hsn sac", "sac"],
    "rate": ["rate", "invoice rate", "unit rate"],
    "quantity": ["quantity", "qty"],
    "unit": ["unit", "uom", "units"],
    "gst": ["gst", "gst rate", "gst percent", "tax", "tax rate"],
    "po_date": ["po date", "purchase order date"],
    "po_number": ["po number", "po no", "po", "purchase order", "purchase order no"],
    "po_rate": ["po rate"],
    "po_amount": ["po amount", "po value"],
}
REQUIRED_IMPORT_FIELDS = ("invoice_date", "invoice_no", "vendor_name", "particulars")


def _norm_header(value):
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").lower()).strip()


ALIAS_LOOKUP = {_norm_header(a): field for field, aliases in IMPORT_ALIASES.items() for a in aliases}


def style_header_row(ws, row=1):
    for cell in ws[row]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 30


def autosize(ws, min_width=8, max_width=45):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            if isinstance(cell.value, (date, datetime)):
                length = 10
            elif isinstance(cell.value, float):
                length = len("{:,.2f}".format(cell.value))
            else:
                length = max(len(line) for line in str(cell.value).split("\n"))
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), length)
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = max(min_width, min(max_width, width + 3))


def to_date(iso):
    return datetime.strptime(iso, "%Y-%m-%d") if iso else None


def xlsx_response(wb, filename):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, mimetype=XLSX_MIME, as_attachment=True, download_name=filename)


@app.get("/api/export")
def api_export():
    filters = read_filters(request.args)
    invoices = fetch_invoices(get_db(), filters)

    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append([c[0] for c in EXPORT_COLUMNS])
    style_header_row(ws)

    for inv in invoices:
        for item in inv["items"] or [{}]:
            values = []
            for _header, key, _fmt in EXPORT_COLUMNS:
                val = inv.get(key) if key in EXPORT_HEADER_KEYS else item.get(key)
                if key in ("invoice_date", "po_date"):
                    val = to_date(val)
                values.append(val)
            ws.append(values)
            r = ws.max_row
            for col, (_header, key, fmt) in enumerate(EXPORT_COLUMNS, start=1):
                cell = ws.cell(row=r, column=col)
                cell.border = BORDER
                if fmt:
                    cell.number_format = fmt
                if key == "po_status" and cell.value in STATUS_FILLS:
                    cell.fill = STATUS_FILLS[cell.value]

    last_row = ws.max_row
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(EXPORT_COLUMNS)), max(last_row, 1))
    autosize(ws)

    # Summary sheet
    ss = wb.create_sheet("Summary")
    items = [it for inv in invoices for it in inv["items"]]
    rng = "%s to %s" % (filters["start_date"] or "beginning", filters["end_date"] or "today")
    rows = [
        ("Report", "Store Invoice & PO Summary"),
        ("Generated", datetime.now().strftime("%d-%m-%Y %H:%M")),
        ("Date range", rng),
        ("Other filters", ", ".join("%s=%s" % (k, v) for k, v in filters.items()
                                    if v and k not in ("start_date", "end_date")) or "None"),
        ("Total invoices", len(invoices)),
        ("Total items", len(items)),
        ("Total quantity", round(sum(i["quantity"] or 0 for i in items), 3)),
        ("Value without GST", round2(sum(i["totals"]["without_gst"] for i in invoices))),
        ("GST amount", round2(sum(i["totals"]["gst_amount"] for i in invoices))),
        ("Total invoice value", round2(sum(i["totals"]["amount"] for i in invoices))),
        ("Total PO value", round2(sum(i["totals"]["po_amount"] for i in invoices))),
        ("Total difference", round2(sum(i["totals"]["difference"] for i in invoices))),
    ] + [("%s items" % st, sum(1 for i in items if i["po_status"] == st)) for st in PO_STATUSES]
    for label, value in rows:
        ss.append([label, value])
        ss.cell(row=ss.max_row, column=1).font = Font(bold=True)
        if isinstance(value, float):
            ss.cell(row=ss.max_row, column=2).number_format = FMT_MONEY
    ss["B1"].font = Font(bold=True, size=13)
    ss.column_dimensions["A"].width = 22
    ss.column_dimensions["B"].width = 40

    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    return xlsx_response(wb, "invoices_%s.xlsx" % stamp)


@app.get("/api/template")
def api_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    ws.append([c[0] for c in TEMPLATE_COLUMNS])
    style_header_row(ws)

    today = date.today()
    d1, d2 = today - timedelta(days=1), today
    samples = [
        [d1, "INV1001", "Sharma Traders", "Purchase", "Pending", "Yes", "Morning delivery",
         "Maida", "1101", 32.20, 5000, "KG", 5, d1, "PO1001", 32.00, 160000],
        [d1, "INV1001", "Sharma Traders", "Purchase", "Pending", "Yes", "Morning delivery",
         "Sugar", "1701", 40.00, 7000, "KG", 5, d1, "PO1002", 40.00, 280000],
        [d1, "INV1001", "Sharma Traders", "Purchase", "Pending", "Yes", "Morning delivery",
         "Salt", "2501", 12.00, 1000, "KG", 5, d1, "PO1003", 11.80, 11800],
        [d2, "BL-552", "Gupta Packaging", "Purchase", "Paid", "Yes", "",
         "Carton Box 12x12", "4819", 18.50, 2000, "NOS", 18, None, "", None, None],
    ]
    for row in samples:
        ws.append([datetime.combine(v, datetime.min.time()) if isinstance(v, date) else v for v in row])
    for r in range(2, 202):  # pre-format rows so typed values look right
        for col, (_h, fmt, _w) in enumerate(TEMPLATE_COLUMNS, start=1):
            if fmt:
                ws.cell(row=r, column=col).number_format = fmt
    for col, (_h, _fmt, width) in enumerate(TEMPLATE_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s1" % get_column_letter(len(TEMPLATE_COLUMNS))

    for col, options in (("D", INVOICE_TYPES), ("E", PAYMENT_STATUSES), ("F", RECEIVED_OPTIONS)):
        dv = DataValidation(type="list", formula1='"%s"' % ",".join(options), allow_blank=True,
                            showErrorMessage=True, errorTitle="Invalid value",
                            error="Choose one of: %s" % ", ".join(options))
        dv.add("%s2:%s1000" % (col, col))
        ws.add_data_validation(dv)

    ins = wb.create_sheet("Instructions")
    lines = [
        ("How to fill the Invoices sheet", None),
        ("", None),
        ("1. One row = one item. An invoice with 3 items needs 3 rows.", None),
        ("2. Rows with the same Invoice Date + Invoice No. + Vendor Name become ONE invoice.", None),
        ("   Example: rows 2-4 (INV1001 / Sharma Traders) are imported as one invoice with 3 items.", None),
        ("3. Do not rename or delete the header row. Extra columns are ignored.", None),
        ("", None),
        ("Column", "Rules"),
        ("Invoice Date", "Required. A real Excel date or text like 01-10-2026 (day-month-year)."),
        ("Invoice No.", "Required."),
        ("Vendor Name", "Required."),
        ("Invoice Type", "Purchase / Service / Other. Blank = Purchase."),
        ("Payment Status", "Pending / Paid / Partially Paid / Hold. Blank = Pending."),
        ("Invoice Received", "Yes / No. Blank = Yes."),
        ("Remarks", "Optional."),
        ("Particulars", "Required. Item / material name."),
        ("HSN", "Optional. HSN / SAC code."),
        ("Rate", "Required. Number, 0 or more. Rate per unit before GST."),
        ("Quantity", "Required. Number, 0 or more."),
        ("Unit", "Optional. KG, NOS, LTR, BAG ..."),
        ("GST %", "Optional number (5 means 5%). Blank = 0."),
        ("PO Date", "Optional date."),
        ("PO Number", "Optional. Leave blank if the PO is not available yet (status = PO Pending)."),
        ("PO Rate", "Optional number."),
        ("PO Amount", "Optional number. May be with or without GST."),
        ("", None),
        ("Calculated automatically (do not enter)", None),
        ("Without GST", "Rate x Quantity"),
        ("Amount", "Rate x Quantity x (1 + GST% / 100)"),
        ("PO Rate Difference", "Rate - PO Rate"),
        ("PO Amount Difference", "Amount - PO Amount"),
        ("PO Match Status", "Matched / Rate Mismatch / PO Pending / Check Rate"),
        ("", None),
        ("Import options", None),
        ("Existing invoices", "An invoice with the same Invoice No. and Vendor that already exists is "
                              "skipped, unless you choose 'Replace it with the data from Excel' when importing."),
        ("Errors", "If any row has a problem nothing is imported, and the row numbers are listed."),
    ]
    for a, b in lines:
        ins.append([a, b])
        if b is None and a:
            ins.cell(row=ins.max_row, column=1).font = Font(bold=True, size=12, color="1F4E78")
        if a == "Column" or a.startswith("Calculated") or a == "Import options":
            for c in ins[ins.max_row]:
                c.font = Font(bold=True)
    ins.column_dimensions["A"].width = 32
    ins.column_dimensions["B"].width = 95
    for row in ins.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    return xlsx_response(wb, "invoice_import_template.xlsx")


def _check_xlsx_upload():
    f = request.files.get("file")
    if f is None or not f.filename:
        raise ApiError("Please choose an Excel (.xlsx) file to import.")
    if not f.filename.lower().endswith((".xlsx", ".xlsm")):
        raise ApiError("Only .xlsx Excel files are supported. In Excel use File > Save As > .xlsx.")
    data = f.read()
    if not data:
        raise ApiError("The uploaded file is empty.")
    if not data.startswith(b"PK"):
        raise ApiError("This file is not a valid .xlsx Excel file.")
    return data


def _find_header_row(ws):
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=15, values_only=True), start=1):
        mapping = {}
        for c_idx, value in enumerate(row):
            field = ALIAS_LOOKUP.get(_norm_header(value))
            if field and field not in mapping:
                mapping[field] = c_idx
        if all(f in mapping for f in ("invoice_no", "particulars")):
            return r_idx, mapping
    return None, None


@app.post("/api/import")
def api_import():
    data = _check_xlsx_upload()
    mode = (request.form.get("mode") or "skip").lower()
    if mode not in ("skip", "replace"):
        raise ApiError("Unknown import mode.")
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    except Exception:
        log.warning("Could not open uploaded workbook", exc_info=True)
        raise ApiError("Could not read the Excel file. Make sure it is a valid .xlsx file.")

    ws = wb["Invoices"] if "Invoices" in wb.sheetnames else wb.worksheets[0]
    header_row, mapping = _find_header_row(ws)
    if header_row is None:
        wb.close()
        raise ApiError("Could not find the header row. The sheet needs at least these columns: "
                       "Invoice Date, Invoice No., Vendor Name, Particulars. "
                       "Download the Excel Template to see the correct format.")
    missing = [f for f in REQUIRED_IMPORT_FIELDS if f not in mapping]
    if missing:
        wb.close()
        names = {"invoice_date": "Invoice Date", "invoice_no": "Invoice No.",
                 "vendor_name": "Vendor Name", "particulars": "Particulars"}
        raise ApiError("Missing required column(s): %s." % ", ".join(names[m] for m in missing))

    groups, order, errors = {}, [], []
    row_count = 0
    for r_idx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True),
                                start=header_row + 1):
        values = {f: (row[c] if c < len(row) else None) for f, c in mapping.items()}
        if all(v is None or (isinstance(v, str) and not v.strip()) for v in values.values()):
            continue
        row_count += 1
        try:
            inv_date = parse_date(values.get("invoice_date"))
        except ValueError:
            errors.append("Row %d: Invoice Date '%s' is not a valid date." % (r_idx, values.get("invoice_date")))
            continue
        inv_no = clean_text(values.get("invoice_no"))
        vendor = clean_text(values.get("vendor_name"))
        if not inv_date or not inv_no or not vendor:
            errors.append("Row %d: Invoice Date, Invoice No. and Vendor Name are required." % r_idx)
            continue
        key = (inv_date, inv_no.lower(), vendor.lower())
        if key not in groups:
            groups[key] = {f: values.get(f) for f in ("invoice_type", "payment_status",
                                                      "invoice_received", "remarks")}
            groups[key].update({"invoice_date": inv_date, "invoice_no": inv_no,
                                "vendor_name": vendor, "items": [], "_first_row": r_idx})
            order.append(key)
        item = {f: values.get(f) for f in ("particulars", "hsn", "rate", "quantity", "unit", "gst",
                                           "po_date", "po_number", "po_rate", "po_amount")}
        item["_label"] = "Row %d" % r_idx
        groups[key]["items"].append(item)
    wb.close()

    if row_count == 0:
        raise ApiError("The Excel sheet has no data rows below the header.")

    validated = []
    for key in order:
        grp = groups[key]
        header, items, errs = validate_invoice(grp)
        for e in errs:
            errors.append(e if e.startswith("Row ") else "Row %d (%s): %s" % (grp["_first_row"],
                                                                           grp["invoice_no"], e))
        validated.append((header, items))
    if errors:
        raise ApiError("Import cancelled. Please fix %d problem(s) in the Excel file and try again."
                       % len(errors), 400, errors[:100])

    db = get_db()
    created = updated = 0
    skipped = []
    with db:
        for header, items in validated:
            existing = find_duplicate(db, header)
            if existing and mode == "skip":
                skipped.append("%s (%s)" % (header["invoice_no"], header["vendor_name"]))
                continue
            if existing:
                update_invoice(db, existing, header, items)
                updated += 1
            else:
                insert_invoice(db, header, items)
                created += 1

    items_total = sum(len(items) for _h, items in validated)
    msg = "Imported %d row(s): %d invoice(s) created" % (row_count, created)
    if updated:
        msg += ", %d replaced" % updated
    if skipped:
        msg += ", %d skipped (already exist)" % len(skipped)
    return jsonify({"message": msg + ".", "rows": row_count, "invoices": len(validated),
                    "items": items_total, "created": created, "updated": updated,
                    "skipped": skipped})


# ---------------------------------------------------------------------------
# OCR
#
# Providers are pluggable. To add one (e.g. a cloud OCR API):
#   1. subclass BaseOCRProvider and implement extract_text(image_bytes, filename)
#   2. read any API key from an environment variable (never hard-code it)
#   3. register it in OCR_PROVIDERS and start the app with OCR_PROVIDER=<name>
# ---------------------------------------------------------------------------

class OCRError(Exception):
    status = 422


class OCRUnavailable(OCRError):
    status = 503


class BaseOCRProvider:
    name = "base"

    def extract_text(self, image_bytes, filename):
        raise NotImplementedError


class TesseractOCRProvider(BaseOCRProvider):
    """Local OCR using the Tesseract engine (installed separately) via pytesseract."""

    name = "tesseract"
    WINDOWS_PATHS = (r"C:\Program Files\Tesseract-OCR\tesseract.exe",
                     r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe")

    def __init__(self):
        try:
            import pytesseract
            from PIL import Image, ImageOps
        except ImportError:
            raise OCRUnavailable("OCR libraries are not installed. Run: pip install Pillow pytesseract")
        self.pytesseract, self.Image, self.ImageOps = pytesseract, Image, ImageOps
        cmd = os.environ.get("TESSERACT_CMD")
        if not cmd and os.name == "nt":
            cmd = next((p for p in self.WINDOWS_PATHS if os.path.exists(p)), None)
        if cmd:
            pytesseract.pytesseract.tesseract_cmd = cmd
        self.lang = os.environ.get("TESSERACT_LANG", "eng")

    def extract_text(self, image_bytes, filename):
        try:
            img = self.Image.open(io.BytesIO(image_bytes))
            img = self.ImageOps.exif_transpose(img)
            img = img.convert("L")
            if img.width < 1600:  # small photos OCR better when upscaled
                scale = 1600 / float(img.width)
                img = img.resize((1600, int(img.height * scale)))
        except Exception:
            raise OCRError("Could not read the image. Please upload a clear JPG or PNG photo.")
        try:
            return self.pytesseract.image_to_string(img, lang=self.lang, config="--psm 6")
        except self.pytesseract.TesseractNotFoundError:
            raise OCRUnavailable(
                "The Tesseract OCR engine is not installed or not found. Install it (see README) "
                "or set the TESSERACT_CMD environment variable to tesseract.exe.")
        except Exception:
            log.exception("Tesseract failed")
            raise OCRError("OCR failed on this image. Try a clearer, straight photo.")


OCR_PROVIDERS = {"tesseract": TesseractOCRProvider}


def get_ocr_provider():
    name = os.environ.get("OCR_PROVIDER", "tesseract").strip().lower()
    cls = OCR_PROVIDERS.get(name)
    if cls is None:
        raise OCRUnavailable("OCR provider '%s' is not configured." % name)
    return cls()


def extract_pdf_text(data):
    """Text-based (digital) PDFs only; scanned PDFs have no text layer."""
    try:
        from pypdf import PdfReader
    except ImportError:
        raise OCRUnavailable("PDF support needs the 'pypdf' package: pip install pypdf")
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = reader.pages[:5]
        return "\n".join((p.extract_text() or "") for p in pages)
    except Exception:
        raise OCRError("Could not read this PDF file.")


UNIT_WORDS = {"KG", "KGS", "GM", "GMS", "G", "QTL", "TON", "TONS", "MT", "LTR", "LTRS", "LT", "L",
              "ML", "NOS", "NO", "PCS", "PC", "BAG", "BAGS", "BOX", "BOXES", "PKT", "PKTS", "CTN",
              "DOZ", "SET", "SETS", "MTR", "MTRS", "ROLL", "ROLLS", "BTL", "DRUM", "UNIT", "UNITS",
              "EA", "EACH", "PAIR", "BUNDLE", "SQM", "SQFT", "CAN", "TIN", "JAR", "PACK"}
NUM_TOKEN = re.compile(r"^(?:rs\.?|inr|\u20b9)?(\d{1,3}(?:,\d{2,3})+(?:\.\d+)?|\d+(?:\.\d+)?)(%?)$",
                       re.IGNORECASE)
DATE_IN_TEXT = re.compile(
    r"\b(\d{1,2}[-/.]\d{1,2}[-/.](?:\d{4}|\d{2})|\d{4}[-/.]\d{1,2}[-/.]\d{1,2}"
    r"|\d{1,2}(?:st|nd|rd|th)?[\s\-/.]*(?:jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*"
    r"[\s\-/.,]*\d{2,4})\b", re.IGNORECASE)
SKIP_ITEM_LINE = re.compile(
    r"\b(total|sub\s*total|grand|cgst|sgst|igst|gstin|gst\s*no|invoice|bill\s*no|date|phone|mobile|"
    r"email|bank|ifsc|a/c|account|round|words|signature|authori[sz]ed|terms|address|"
    r"freight|discount|balance|due|vehicle|transport|qty\b.*rate|description|particulars)\b",
    re.IGNORECASE)
COMPANY_WORDS = re.compile(
    r"\b(pvt|private|ltd|limited|llp|traders?|trading|enterprises?|industries|corporation|corp|"
    r"company|agenc(?:y|ies)|stores?|suppliers?|foods?|mills?|distributors?|marketing|sons|brothers|"
    r"impex|exports?|mart|udyog)\b", re.IGNORECASE)


def _num(token):
    m = NUM_TOKEN.match(token.strip(",:;()"))
    if not m:
        return None, False
    return float(m.group(1).replace(",", "")), bool(m.group(2))


def _close(a, b, rel=0.01):
    return abs(a - b) <= max(1.0, abs(b) * rel)


GST_SLABS = (3.0, 5.0, 12.0, 18.0, 28.0)


def _find_qty_rate_amount(numbers, gst, qty_hint):
    """Find quantity x rate (x GST) = amount among the numbers of one line."""
    for k in range(len(numbers) - 1, 1, -1):
        for i in range(k):
            for j in range(k):
                if i == j or not numbers[i] * numbers[j]:
                    continue
                prod = numbers[i] * numbers[j]
                candidates = [gst] if gst is not None else [0.0] + [
                    n for x, n in enumerate(numbers) if n in GST_SLABS and x not in (i, j, k)]
                for pct in candidates:
                    if _close(prod * (1 + pct / 100.0), numbers[k]):
                        q, r = (i, j) if (qty_hint == i or (qty_hint != j and i < j)) else (j, i)
                        return numbers[q], numbers[r], numbers[k], pct
    return None


def _parse_item_line(line):
    if SKIP_ITEM_LINE.search(line):
        return None
    tokens = [t for t in re.split(r"[\s|]+", line.strip()) if t]
    if len(tokens) < 3:
        return None
    if _num(tokens[0])[0] is not None and len(tokens[0]) <= 3 and "." not in tokens[0] \
            and _num(tokens[1])[0] is None:
        tokens = tokens[1:]  # leading serial number
    desc = []
    while tokens and _num(tokens[0])[0] is None:
        desc.append(tokens.pop(0))
    particulars = " ".join(desc).strip(" -:.,")
    if len(re.sub(r"[^A-Za-z]", "", particulars)) < 2:
        return None
    raw, unit, gst, qty_hint = [], "", None, None  # raw = [(value, token)]
    for tok in tokens:
        up = tok.upper().strip(".,")
        if up in UNIT_WORDS:
            unit = up
            if raw and qty_hint is None:
                qty_hint = len(raw) - 1  # the number just before the unit is usually the quantity
            continue
        value, is_pct = _num(tok)
        if value is None:
            continue
        if is_pct:
            gst = value
        else:
            raw.append((value, tok))
    numbers = [v for v, _t in raw]

    # A leading 4-8 digit integer is probably the HSN code - unless the unit follows it (a quantity).
    hsn_candidate = bool(raw) and re.fullmatch(r"\d{4,8}", raw[0][1]) is not None and qty_hint != 0
    hsn, match = "", None
    if hsn_candidate:
        match = _find_qty_rate_amount(numbers[1:], gst, None if qty_hint is None else qty_hint - 1)
        if match:
            hsn = raw[0][1]
    if match is None:
        match = _find_qty_rate_amount(numbers, gst, qty_hint)
    if match is None and hsn_candidate and len(numbers) >= 3:
        hsn, numbers = raw[0][1], numbers[1:]
        qty_hint = None if qty_hint is None else qty_hint - 1
    if match is None and len(numbers) < 2:
        return None

    amount = None
    if match:
        qty, rate, amount, found_gst = match
        if gst is None and found_gst:
            gst = found_gst
    elif qty_hint is not None and qty_hint + 1 < len(numbers):
        qty, rate = numbers[qty_hint], numbers[qty_hint + 1]
    else:
        qty, rate = numbers[0], numbers[1]
    return {"particulars": particulars, "hsn": hsn, "rate": rate, "quantity": qty, "unit": unit,
            "gst": gst, "amount": amount}


def parse_invoice_text(text):
    """Best-effort extraction of invoice fields from OCR text."""
    lines = [re.sub(r"[ \t]+", " ", ln).strip() for ln in text.splitlines()]
    lines = [ln for ln in lines if ln]
    result = {"invoice_date": None, "invoice_no": "", "vendor_name": "", "po_number": "",
              "invoice_amount": None, "items": []}

    for m in re.finditer(r"\b(?:invoice|inv|bill)\s*(?:no\b|num\b|number\b|#)\.?\s*[:#.\-]*\s*"
                         r"([A-Za-z0-9][A-Za-z0-9\-/]{0,30})", text, re.IGNORECASE):
        if re.search(r"\d", m.group(1)):
            result["invoice_no"] = m.group(1).strip("-/")
            break

    for ln in lines:  # prefer a date on a line labelled "date" but not "PO date" / "due date"
        if re.search(r"date", ln, re.IGNORECASE) and not re.search(r"\b(po|due|order|dc)\b.{0,6}date", ln,
                                                                   re.IGNORECASE):
            dm = DATE_IN_TEXT.search(ln)
            if dm:
                try:
                    result["invoice_date"] = parse_date(dm.group(1))
                    break
                except ValueError:
                    pass
    if not result["invoice_date"]:
        for dm in DATE_IN_TEXT.finditer(text):
            try:
                result["invoice_date"] = parse_date(dm.group(1))
                break
            except ValueError:
                continue

    for m in re.finditer(r"(?:\bP\.?\s?O\b\.?|\bpurchase\s*order|\border)\s*"
                         r"(?:no\b|number\b|#|ref\b)?\.?\s*[:#.\-]*\s*"
                         r"([A-Za-z0-9][A-Za-z0-9\-/]{0,30})", text, re.IGNORECASE):
        if re.search(r"\d", m.group(1)):
            result["po_number"] = m.group(1).strip("-/")
            break
    if not result["po_number"]:
        m = re.search(r"\bPO[\-/]?\d{2,}\b", text, re.IGNORECASE)
        if m:
            result["po_number"] = m.group(0)

    m = re.search(r"(?:vendor|supplier|seller|sold\s*by|billed\s*by|from)\s*(?:name)?\s*[:\-]\s*(.+)",
                  text, re.IGNORECASE)
    if m:
        result["vendor_name"] = m.group(1).strip()
    else:
        for ln in lines[:12]:
            if COMPANY_WORDS.search(ln) and not re.search(r"bill\s*to|ship\s*to|buyer|consignee", ln,
                                                          re.IGNORECASE):
                result["vendor_name"] = ln
                break
        if not result["vendor_name"]:
            for ln in lines[:6]:
                if len(re.sub(r"[^A-Za-z]", "", ln)) >= 4 and not re.search(
                        r"invoice|original|duplicate|gstin|copy|cash|credit|memo", ln, re.IGNORECASE):
                    result["vendor_name"] = ln
                    break
    result["vendor_name"] = re.sub(r"^m/?s\.?\s*", "", result["vendor_name"], flags=re.IGNORECASE)[:150]

    for label in (r"grand\s*total", r"invoice\s*(?:total|value|amount)", r"net\s*(?:amount|payable|total)",
                  r"total\s*(?:amount|payable|value)", r"amount\s*payable", r"total"):
        found = re.findall(label + r"[^0-9\n]{0,25}(\d{1,3}(?:,\d{2,3})+(?:\.\d+)?|\d+(?:\.\d+)?)",
                           text, re.IGNORECASE)
        if found:
            result["invoice_amount"] = max(float(f.replace(",", "")) for f in found)
            break

    for ln in lines:
        item = _parse_item_line(ln)
        if item:
            result["items"].append(item)
        if len(result["items"]) >= 100:
            break
    if result["po_number"]:
        for item in result["items"]:
            item["po_number"] = result["po_number"]
    return result


IMAGE_SIGNATURES = {".jpg": b"\xff\xd8\xff", ".jpeg": b"\xff\xd8\xff", ".png": b"\x89PNG"}


@app.post("/api/ocr")
def api_ocr():
    f = request.files.get("file")
    if f is None or not f.filename:
        raise ApiError("Please choose an invoice image (JPG / PNG) or PDF.")
    ext = os.path.splitext(f.filename.lower())[1]
    if ext not in (".jpg", ".jpeg", ".png", ".pdf"):
        raise ApiError("Unsupported file type. Upload a JPG, JPEG, PNG or PDF file.")
    data = f.read()
    if not data:
        raise ApiError("The uploaded file is empty.")

    try:
        if ext == ".pdf":
            if not data.startswith(b"%PDF"):
                raise ApiError("This file is not a valid PDF.")
            text, provider = extract_pdf_text(data), "pdf-text"
            if not text.strip():
                raise OCRError("This PDF has no readable text (it is probably a scanned image). "
                               "Take a photo or screenshot of the invoice and upload it as JPG/PNG.")
        else:
            if not data.startswith(IMAGE_SIGNATURES[ext]):
                raise ApiError("The file content does not match its extension (%s)." % ext)
            ocr = get_ocr_provider()
            text, provider = ocr.extract_text(data, f.filename), ocr.name
    except OCRError as err:
        return jsonify({"error": str(err)}), err.status

    parsed = parse_invoice_text(text or "")
    found = [k for k in ("invoice_date", "invoice_no", "vendor_name", "po_number", "invoice_amount")
             if parsed.get(k)]
    return jsonify({
        "provider": provider,
        "extracted": parsed,
        "fields_found": found,
        "items_found": len(parsed["items"]),
        "raw_text": (text or "")[:20000],
        "message": "OCR finished. Please verify every highlighted value before saving.",
    })


# ---------------------------------------------------------------------------

init_db()
